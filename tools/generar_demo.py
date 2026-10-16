"""Genera los Excels ficticios de demostración de Pingüino Delta (semilla fija = reproducible).
Uso: python3 tools/generar_demo.py   (requiere openpyxl)
"""
import json, random, math, re, os, datetime as dt
BASE=os.path.dirname(os.path.abspath(__file__))
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
random.seed(2026)
_h=open(os.path.join(BASE,'..','index.html'),encoding='utf-8').read()
ARTS=set(json.loads(re.search(r'const ARTS=(\[.*?\]);',_h).group(1)))
MESES=["Enero","Febrero","Marzo","Abril","Mayo","Junio","Julio","Agosto","Septiembre","Octubre","Noviembre","Diciembre"]
HOY=dt.date(2026,10,16)
# (artículo, consumo base mensual, estacional respiratorio, escenario, costo unitario CLP, origen)
D=[
("LOSARTAN (POTASICO) 50 MG COMPRIMIDO",24000,0,'ok',9,'CENABAST'),
("METFORMINA (CLORHIDRATO) 850 MG COMPRIMIDO RECUBIERTO",31000,0,'ok',14,'CENABAST'),
("METFORMINA (CLORHIDRATO) 1 GR COMPRIMIDO DE LIBERACIÓN PROLONGADA",6500,0,'quiebre',62,'CENABAST'),
("ENALAPRIL (MALEATO) 10 MG COMPRIMIDO",5200,0,'sobrestock',6,'CENABAST'),
("ATORVASTATINA 20 MG COMPRIMIDO",15500,0,'ok',18,'CENABAST'),
("AMLODIPINO 5 MG COMPRIMIDO",9800,0,'merma',7,'CENABAST'),
("AMLODIPINO 10 MG COMPRIMIDO",4100,0,'ok',9,'CENABAST'),
("ACIDO ACETILSALICILICO 100 MG COMPRIMIDO",8700,0,'ok',5,'CENABAST'),
("HIDROCLOROTIAZIDA 50 MG COMPRIMIDO",2300,0,'sobrestock',11,'CENABAST'),
("CARVEDILOL 12.5 MG COMPRIMIDO",3600,0,'ok',21,'CENABAST'),
("CARVEDILOL 25 MG COMPRIMIDO",1900,0,'quiebre',29,'CENABAST'),
("ESPIRONOLACTONA 25 MG COMPRIMIDO",1400,0,'ok',24,'CENABAST'),
("FUROSEMIDA 40 MG COMPRIMIDO",2100,0,'merma',8,'CENABAST'),
("GLIBENCLAMIDA 5 MG COMPRIMIDO",1700,0,'sobrestock',6,'CENABAST'),
("INSULINA NPH (ISOFANA) 100 UI/ML SOLUCIÓN INYECTABLE FRASCO AMPOLLA",210,0,'ok',3450,'SSVQP'),
("JERINGA 100 UI INSULINA",3900,0,'ok',95,'SSVQP'),
("LEVOTIROXINA SODICA 100 MCG COMPRIMIDO",5400,0,'ok',19,'CENABAST'),
("LEVOTIROXINA SODICA 50 MCG COMPRIMIDO",4800,0,'quiebre',16,'CENABAST'),
("OMEPRAZOL 20 MG CAPSULA",9100,0,'ok',12,'CENABAST'),
("SERTRALINA 50 MG COMPRIMIDO",6200,0,'ok',23,'CENABAST'),
("FLUOXETINA 20 MG COMPRIMIDO",4300,0,'merma',15,'CENABAST'),
("QUETIAPINA 25 MG COMPRIMIDO RECUBIERTO",2600,0,'ok',38,'CENABAST'),
("CLONAZEPAM 2 MG COMPRIMIDO",1800,0,'ok',10,'CENABAST'),
("PARACETAMOL 500 MG COMPRIMIDO",12500,1,'ok',7,'CENABAST'),
("PARACETAMOL 100 MG/ML SOLUCIÓN GOTAS ORALES",320,1,'quiebre',420,'CENABAST'),
("IBUPROFENO 400 MG COMPRIMIDO",7400,0,'ok',11,'CENABAST'),
("IBUPROFENO 200 MG/5 ML JARABE",260,1,'merma',690,'CENABAST'),
("SALBUTAMOL AEROSOL 100 MCG/DOSIS INHALADOR",540,1,'quiebre',1480,'CENABAST'),
("AMOXICILINA 500 MG COMPRIMIDO",5100,1,'ok',27,'CENABAST'),
("AMOXICILINA 500 MG/5 ML POLVO PARA SUSPENSIÓN ORAL",190,1,'ok',890,'CENABAST'),
("AMOXICILINA + ACIDO CLAVULANICO 875 MG/125 MG COMPRIMIDO",1600,1,'sobrestock',120,'CENABAST'),
("AZITROMICINA 500 MG COMPRIMIDO RECUBIERTO",700,1,'ok',140,'CENABAST'),
("CEFADROXILO 500 MG COMPRIMIDO RECUBIERTO",1100,0,'ok',85,'CENABAST'),
("NITROFURANTOINA 100 MG CAPSULA",1500,0,'merma',48,'CENABAST'),
("METRONIDAZOL 500 MG COMPRIMIDO",900,0,'ok',33,'CENABAST'),
("PREDNISONA 20 MG COMPRIMIDO",1300,1,'ok',22,'CENABAST'),
("CLORFENAMINA 4 MG COMPRIMIDO",2200,1,'sobrestock',4,'CENABAST'),
("LORATADINA 5 MG/5 ML JARABE",140,1,'ok',560,'CENABAST'),
("DOMPERIDONA 10 MG COMPRIMIDO",1200,0,'ok',17,'CENABAST'),
("SALES DE REHIDRATACION ORAL 60 MEQ DE SODIO/LITRO",380,0,'merma',210,'CENABAST'),
]
for a,*_ in D: assert a in ARTS, a
# Doble programación (CENABAST + SSVQP) para mostrar desglose por origen
DOBLE={"LOSARTAN (POTASICO) 50 MG COMPRIMIDO":0.25,"SALBUTAMOL AEROSOL 100 MCG/DOSIS INHALADOR":0.30,"METFORMINA (CLORHIDRATO) 850 MG COMPRIMIDO RECUBIERTO":0.20}

def est(m,flag):  # m 1..12 ; pico invernal junio-agosto
    return 1+0.45*math.cos((m-7)/12*2*math.pi) if flag else 1
def r(x): return int(round(x))

# ---------- consumo mes a mes ----------
cons={}  # art -> year -> [12]
for a,b,e,sc,c,o in D:
    cons[a]={}
    for y,tr in ((2024,0.90),(2025,0.96),(2026,1.0)):
        fila=[]
        for m in range(1,13):
            if y==2026 and m>9: fila.append(None); continue
            v=b*tr*est(m,e)*random.uniform(0.88,1.12)
            fila.append(r(v))
        cons[a][y]=fila
prom={a:r(sum(v for v in cons[a][2026] if v)/9) for a in cons}

# ---------- estilos ----------
NAVY="1F3864"; hdrF=Font(bold=True,color="FFFFFF",name="Calibri"); hdrP=PatternFill("solid",fgColor=NAVY)
thin=Side(style="thin",color="BFBFBF")
def hoja(ws,headers,rows,widths=None,datecols=(),numfmt=None):
    ws.append(headers)
    for i,_ in enumerate(headers,1):
        cl=ws.cell(1,i); cl.font=hdrF; cl.fill=hdrP; cl.alignment=Alignment(horizontal="center",vertical="center",wrap_text=True)
    for rw in rows: ws.append(rw)
    for row in ws.iter_rows(min_row=2):
        for cl in row:
            cl.border=Border(bottom=thin)
            if cl.column-1 in datecols: cl.number_format="DD/MM/YYYY"
            elif isinstance(cl.value,(int,float)): cl.number_format=numfmt.get(cl.column-1,"#,##0") if numfmt else "#,##0"
    for i,h in enumerate(headers,1):
        ws.column_dimensions[get_column_letter(i)].width=(widths[i-1] if widths else max(12,len(str(h))+2))
    ws.freeze_panes="B2"; ws.row_dimensions[1].height=30
def guardar(nombre,headers,rows,**kw):
    wb=Workbook(); ws=wb.active; ws.title=kw.pop('titulo','Datos'); hoja(ws,headers,rows,**kw); wb.save(os.path.join(BASE,'..','demo',nombre))

# 01 Inventario de lotes
inv=[]; 
def lote(): return "L%s%02d%s"%(random.choice("ABCDEFGHJK"),random.randint(10,99),random.randint(100,999))
def venc(meses): 
    d=HOY+dt.timedelta(days=int(meses*30.4)); return dt.date(d.year,d.month,random.choice([28,30]) if d.month!=2 else 28)
stock_total={}
for a,b,e,sc,c,o in D:
    p=prom[a]
    if sc=='ok': lots=[(p*random.uniform(1.4,2.2),random.uniform(14,26)),(p*random.uniform(0.8,1.4),random.uniform(8,13))]
    elif sc=='quiebre': lots=[(p*random.uniform(0.25,0.6),random.uniform(10,20))]
    elif sc=='sobrestock': lots=[(p*random.uniform(5,7),random.uniform(20,30)),(p*random.uniform(3,5),random.uniform(15,22))]
    else: lots=[(p*random.uniform(2.5,4),random.uniform(1.2,2.6)),(p*random.uniform(1.5,2.5),random.uniform(16,24))]
    for q,mv in lots:
        inv.append([a,lote(),venc(mv),r(q)])
    stock_total[a]=sum(r(q) for q,_ in lots)
guardar("01_Inventario_lotes_DEMO.xlsx",["Artículo","Lote","Vencimiento","Stock"],inv,widths=[70,14,14,12],datecols=(2,),titulo="Inventario")

# 02 Consumo promedio
guardar("02_Consumo_promedio_DEMO.xlsx",["Artículo","Consumo promedio mensual"],[[a,prom[a]] for a,*_ in D],widths=[70,18],titulo="Consumo promedio")

# 03 Programación vigente (2026 y 2027)
prog=[]
for a,b,e,sc,c,o in D:
    p=prom[a]
    f={'ok':1.0,'quiebre':0.75,'sobrestock':1.25,'merma':1.05}[sc]
    for y,tr in ((2026,1.0),(2027,1.04)):
        base=[p*f*tr*est(m,e)/ (1 if not e else 1) for m in range(1,13)]
        if a in DOBLE:
            s=DOBLE[a]
            prog.append([a,o,y]+[r(v*(1-s)) for v in base])
            prog.append([a,'SSVQP',y]+[r(v*s) for v in base])
        else:
            prog.append([a,o,y]+[r(v) for v in base])
guardar("03_Programacion_vigente_DEMO.xlsx",["Artículo","Origen","Año"]+MESES,prog,widths=[70,12,8]+[11]*12,numfmt={2:"0"},titulo="Programación")

# 04 Consumo mes a mes (2024-2026)
cm=[]
for a,*_ in D:
    for y in (2024,2025,2026): cm.append([a,y]+cons[a][y])
guardar("04_Consumo_mes_a_mes_DEMO.xlsx",["Artículo","Año"]+MESES,cm,widths=[70,8]+[11]*12,numfmt={1:"0"},titulo="Consumo mensual")

# 05 Archivo único (Origen: Consumo AAAA / CENABAST AAAA / SSVQP AAAA)
un=[]
for a,*_ in D:
    for y in (2025,2026): un.append([a,"Consumo %d"%y,y]+cons[a][y])
for rw in prog: un.append([rw[0],"%s %d"%(rw[1],rw[2]),rw[2]]+rw[3:])
guardar("05_Archivo_unico_Consumo_y_Programacion_DEMO.xlsx",["Artículo","Origen","Año"]+MESES,un,widths=[70,16,8]+[11]*12,numfmt={2:"0"},titulo="Consolidado")

# 06 Despachos CENABAST (calendario) oct-2026 a mar-2027
des=[]
for a,b,e,sc,c,o in D:
    if o!='CENABAST': continue
    dia=random.choice([5,8,12,15,19,22,26])
    for k in range(6):
        y,m=(2026,10+k) if k<3 else (2027,k-2)
        d=min(28,max(1,dia+random.randint(-3,3)))
        if sc=='quiebre' and k==1: continue  # despacho perdido para mostrar brecha
        des.append([a,dt.date(y,m,d)])
des.sort(key=lambda x:(x[1],x[0]))
guardar("06_Despachos_CENABAST_DEMO.xlsx",["Artículo","Fecha despacho"],des,widths=[70,16],datecols=(1,),titulo="Despachos")

# 07 Valor unitario
guardar("07_Valor_unitario_DEMO.xlsx",["Artículo","Costo unitario (CLP, neto)"],[[a,c] for a,b,e,sc,c,o in D],widths=[70,20],titulo="Valor unitario")

# 08 Diferencias de inventario (conteo físico)
dif=[]
for a,b,e,sc,c,o in random.sample(D,14):
    v=r(prom[a]*random.uniform(0.005,0.04))*random.choice([-1,-1,1])
    if v: dif.append([a,v,"Conteo físico 30/09/2026"])
guardar("08_Diferencias_inventario_DEMO.xlsx",["Artículo","Diferencia","Observación"],dif,widths=[70,12,28],numfmt={1:"+#,##0;-#,##0"},titulo="Diferencias")

# 09 Gasto/consumo externo (Artículo · Mes · Año · Cantidad · Costo unit.)
g=[]
for a,b,e,sc,c,o in D:
    for y in (2025,2026):
        for m in range(1,13):
            q=cons[a][y][m-1]
            if q: g.append([a,m,y,q,c])
guardar("09_Gasto_consumo_externo_DEMO.xlsx",["Artículo","Mes","Año","Cantidad","Costo unitario"],g,widths=[70,7,8,12,14],numfmt={1:"0",2:"0"},titulo="Gasto")

print('Archivos 01-09 regenerados en demo/ (00_Guion se mantiene aparte).')
